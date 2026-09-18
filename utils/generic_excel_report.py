"""Generic Excel report generation for collected product workspaces."""

from __future__ import annotations

import json
from collections import OrderedDict
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from models.schemas import Product
from utils.generic_excel_headers import display_column_titles, has_custom_column_titles, write_column_schema_sheet
from utils.product_report_display_fields import distinct_product_brand
from utils.report_filter_facets import alcohol_type as report_alcohol_type, supplier as report_supplier

BASE_COLUMNS = (
    "№|Категория|Мини-каталог|Товар|Бренд|Поставщик/производитель|Цена|"
    "Старая цена|Цена за единицу|В наличии|Вес|Объём|Жирность|Состояние|"
    "Форма товара|Упаковка|Страна|Белки|Жиры|Углеводы|Калорийность|"
    "Состав|Условия хранения|Срок годности|Ссылка|Изображение|ID товара|"
    "Экспортировано"
).split("|")

RAW_FIELD_TITLES = {
    "product_type": "Мини-каталог",
    "brand": "Бренд",
    "supplier": "Поставщик",
    "producer": "Производитель",
    "manufacturer": "Изготовитель",
    "vendor": "Поставщик",
    "country": "Страна",
    "country_of_origin": "Страна происхождения",
    "origin_country": "Страна происхождения",
    "weight": "Вес",
    "volume": "Объём",
    "unit": "Единица",
    "fat": "Жиры",
    "fat_percent": "Жирность",
    "sugar": "Сахар",
    "sugar_class": "Тип сахара",
    "color": "Цвет",
    "variant": "Вариант",
    "package_size": "Размер упаковки",
    "product_state": "Состояние",
    "product_form": "Форма товара",
    "packaging": "Упаковка",
    "protein": "Белки",
    "carbohydrate": "Углеводы",
    "calories": "Калорийность",
    "composition": "Состав",
    "description": "Описание",
    "storage_conditions": "Условия хранения",
    "shelf_life": "Срок годности",
    "special_offers": "Акции",
    "alcohol_type": "Алкогольный тип",
}

CORE_RAW_FIELDS = {
    "product_type",
    "brand",
    "supplier",
    "producer",
    "manufacturer",
    "vendor",
    "country",
    "country_of_origin",
    "origin_country",
    "weight",
    "volume",
    "fat",
    "fat_percent",
    "product_state",
    "product_form",
    "packaging",
    "protein",
    "carbohydrate",
    "calories",
    "composition",
    "storage_conditions",
    "shelf_life",
}

TECHNICAL_RAW_FIELDS = {
    "categories",
    "field_sources",
    "source",
    "source_id",
    "url",
    "link",
    "product_link",
    "image",
    "images",
    "image_url",
}


def write_generic_products_excel_report(
    products: list[Product],
    *,
    shop: str,
    output_dir: Path | str,
    exported_at: str = "",
    selected_columns: list[str] | None = None,
    column_titles: dict[str, str] | None = None,
    products_sheet_title: str = "",
) -> Path:
    """Write collected products into one user-facing Excel table."""
    target_dir = Path(output_dir)
    target_dir.mkdir(parents=True, exist_ok=True)
    report_path = target_dir / f"{shop}_products.xlsx"
    columns = _select_columns(_build_columns(products, exported_at=exported_at), selected_columns)

    workbook = Workbook()
    products_sheet = workbook.active
    products_sheet.title = _safe_sheet_title(products_sheet_title or "Товары")
    _write_products_sheet(products_sheet, products=products, columns=columns, exported_at=exported_at, column_titles=column_titles)

    summary_sheet = workbook.create_sheet(title="Сводка")
    _write_summary_sheet(summary_sheet, shop=shop, products=products, exported_at=exported_at)
    if has_custom_column_titles(columns, column_titles):
        schema_sheet = workbook.create_sheet(title="Столбцы")
        write_column_schema_sheet(schema_sheet, columns=columns, column_titles=column_titles)
    workbook.save(report_path)
    return report_path


def _build_columns(products: list[Product], *, exported_at: str = "") -> list[str]:
    available: OrderedDict[str, None] = OrderedDict()
    rows = [_build_row(product, row_number=index, exported_at=exported_at) for index, product in enumerate(products, 1)]
    default_columns = ["№", "Категория", "Мини-каталог", "Товар", "Цена", "В наличии", "Ссылка", "ID товара"]
    for column in default_columns:
        available[column] = None
    for column in BASE_COLUMNS:
        if column in available:
            continue
        if column == "Экспортировано" and not exported_at:
            continue
        if any(_has_value(row.get(column)) for row in rows):
            available[column] = None
    for product in products:
        raw = dict(product.raw_data or {})
        for key in sorted(raw):
            if key in CORE_RAW_FIELDS or key in TECHNICAL_RAW_FIELDS:
                continue
            value = raw.get(key)
            if not _has_value(value):
                continue
            title = _field_title(key)
            if title not in available:
                available[title] = None
    return list(available)


def _select_columns(available: list[str], selected: list[str] | None) -> list[str]:
    if selected is None:
        return available
    if not selected:
        raise ValueError("No report columns selected")
    chosen = [column for column in selected if column in available]
    if not chosen:
        raise ValueError("Selected report columns are not available")
    return chosen


def _write_summary_sheet(sheet: Worksheet, *, shop: str, products: list[Product], exported_at: str) -> None:
    categories = _category_counts(products)
    rows = [
        ("Магазин", shop),
        ("Товаров", len(products)),
        ("Категорий", len(categories)),
        ("Категории", ", ".join(categories.keys())),
        ("Экспортировано", exported_at),
        ("Шаблон", "Универсальный отчёт по товарам"),
    ]
    for row_index, (label, value) in enumerate(rows, start=1):
        sheet.cell(row=row_index, column=1, value=label).font = Font(bold=True)
        sheet.cell(row=row_index, column=2, value=value)
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 90


def _write_products_sheet(
    sheet: Worksheet,
    *,
    products: list[Product],
    columns: list[str],
    exported_at: str,
    column_titles: dict[str, str] | None = None,
) -> None:
    sheet.append(display_column_titles(columns, column_titles))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for index, product in enumerate(products, start=1):
        row = _build_row(product, row_number=index, exported_at=exported_at)
        sheet.append([row.get(column, "") for column in columns])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    _apply_column_widths(sheet)


def _build_row(product: Product, *, row_number: int, exported_at: str) -> dict[str, Any]:
    raw = dict(product.raw_data or {})
    row: dict[str, Any] = {column: "" for column in BASE_COLUMNS}
    row.update(
        {
            "№": row_number,
            "Категория": str(product.category or ""),
            "Мини-каталог": _first_text(raw, ("product_type",)) or str(product.subcategory or ""),
            "Товар": product.name,
            "Бренд": distinct_product_brand(product),
            "Поставщик/производитель": report_supplier(product),
            "Цена": float(product.price.current),
            "Старая цена": product.price.old if product.price.old is not None else "",
            "Цена за единицу": product.price.unit if product.price.unit is not None else "",
            "В наличии": "Да" if product.in_stock else "Нет",
            "Вес": _first_text(raw, ("weight",)),
            "Объём": _first_text(raw, ("volume",)),
            "Жирность": _first_text(raw, ("fat_percent", "fat")),
            "Состояние": _first_text(raw, ("product_state",)),
            "Форма товара": _first_text(raw, ("product_form",)),
            "Упаковка": _first_text(raw, ("packaging",)),
            "Страна": _first_text(raw, ("country", "country_of_origin", "origin_country")),
            "Белки": _first_text(raw, ("protein",)),
            "Жиры": _first_text(raw, ("fat",)),
            "Углеводы": _first_text(raw, ("carbohydrate",)),
            "Калорийность": _first_text(raw, ("calories",)),
            "Состав": _first_text(raw, ("composition", "description")),
            "Условия хранения": _first_text(raw, ("storage_conditions",)),
            "Срок годности": _first_text(raw, ("shelf_life",)),
            "Ссылка": str(product.product_link),
            "Изображение": str(product.image_url or ""),
            "ID товара": str(product.id or ""),
            "Экспортировано": exported_at,
        }
    )
    for key, value in raw.items():
        if key in CORE_RAW_FIELDS or key in TECHNICAL_RAW_FIELDS or not _has_value(value):
            continue
        row[_field_title(str(key))] = report_alcohol_type(product) if key == "alcohol_type" else _stringify(value)
    return row
def _category_counts(products: list[Product]) -> OrderedDict[str, int]:
    counts: OrderedDict[str, int] = OrderedDict()
    for product in products:
        category_name = str(product.category or "Без категории")
        counts[category_name] = counts.get(category_name, 0) + 1
    return counts
def _field_title(key: str) -> str:
    text = str(key or "").strip()
    if text in RAW_FIELD_TITLES:
        return RAW_FIELD_TITLES[text]
    return "Поле: " + text.replace("_", " ")


def _first_text(raw_data: dict[str, Any], keys: tuple[str, ...]) -> str:
    for key in keys:
        text = _stringify(raw_data.get(key)).strip()
        if text:
            return text
    return ""
def _stringify(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (str, int, float, bool)):
        return str(value)
    return json.dumps(value, ensure_ascii=False, sort_keys=True)


def _has_value(value: Any) -> bool:
    return value not in ("", [], {}, None)
def _apply_column_widths(sheet: Worksheet) -> None:
    width_by_title = {
        "№": 8,
        "Категория": 18,
        "Мини-каталог": 20,
        "Товар": 52,
        "Бренд": 22,
        "Поставщик/производитель": 28,
        "Цена": 12,
        "В наличии": 12,
        "Состав": 56,
        "Ссылка": 60,
        "Изображение": 60,
    }
    for index, cell in enumerate(sheet[1], start=1):
        width = width_by_title.get(str(cell.value), 18)
        sheet.column_dimensions[get_column_letter(index)].width = width


def _safe_sheet_title(value: str) -> str:
    title = "".join("_" if char in r'[]:*?/\\' else char for char in str(value or "").strip())
    return (title or "Товары")[:31]
